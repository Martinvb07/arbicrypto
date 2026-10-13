"""Libro de Excel para cuadrar con la DIAN: resumen explicado, formulario 210, libros y comprobaciones.

Todo va con valores (no fórmulas) para que se vea igual en Excel, Google Sheets o una vista previa."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tax import CO_TZ

COP = '"$"#,##0;[Red]-"$"#,##0;"$"0'
QTY = '#,##0.00######'
DATE = 'dd/mm/yyyy hh:mm'
DAY = 'dd/mm/yyyy'

INK, MUTED, BRAND = "111827", "6B7280", "B7791F"
HEAD = PatternFill("solid", fgColor="1F2937")
SOFT = PatternFill("solid", fgColor="F3F4F6")
SECTION = PatternFill("solid", fgColor="FEF3C7")
GROUPS = {
    "in": (PatternFill("solid", fgColor="DCFCE7"), "166534"),
    "out": (PatternFill("solid", fgColor="FEE2E2"), "991B1B"),
    "bal": (PatternFill("solid", fgColor="DBEAFE"), "1E40AF"),
    "res": (PatternFill("solid", fgColor="FEF3C7"), "92400E"),
    "cost": (PatternFill("solid", fgColor="EDE9FE"), "5B21B6"),
}
OK_FILL, BAD_FILL = PatternFill("solid", fgColor="DCFCE7"), PatternFill("solid", fgColor="FEE2E2")
LINE = Side(style="thin", color="E5E7EB")
STRONG = Side(style="medium", color="9CA3AF")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

KIND = {"compra": "Compra", "venta": "Venta", "permuta": "Cambio entre cripto", "ingreso": "Rendimiento",
        "airdrop": "Airdrop", "comision": "Comisión", "gasto": "Gasto", "aporte": "Aporte de capital", "retiro": "Retiro"}


class F:
    """Valor con un formato distinto al de su columna."""
    def __init__(self, value, fmt):
        self.value, self.fmt = value, fmt


def xl_time(t):
    """Excel no entiende zonas horarias: hora de Colombia sin zona."""
    return datetime.fromtimestamp(t, CO_TZ).replace(tzinfo=None) if t else None


class Sheet:
    def __init__(self, wb, title, widths, first=False):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.fitToWidth, self.ws.page_setup.fitToHeight = 1, 0
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.ncols = len(widths)
        self.row = 1

    def cell(self, col, value, fmt=None, bold=False, color=None, fill=None, align=None, size=None, italic=False):
        c = self.ws.cell(self.row, col, value)
        if fmt:
            c.number_format = fmt
        if bold or color or size or italic:
            c.font = Font(bold=bold, color=color or INK, size=size or 11, italic=italic)
        if fill:
            c.fill = fill
        if align:
            c.alignment = align
        return c

    def skip(self, n=1):
        self.row += n

    def title(self, text, sub=None):
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncols)
        self.cell(1, text, bold=True, size=16)
        self.ws.row_dimensions[self.row].height = 24
        self.row += 1
        if sub:
            self.text(sub, color=MUTED)
        self.row += 1

    def text(self, text, color=None, bold=False, italic=False, height=None):
        """Párrafo que ocupa todo el ancho."""
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncols)
        self.cell(1, text, color=color, bold=bold, italic=italic, align=WRAP)
        if height:
            self.ws.row_dimensions[self.row].height = height
        elif len(str(text)) > 110:
            self.ws.row_dimensions[self.row].height = 15 * (1 + len(str(text)) // 110)
        self.row += 1

    def section(self, text):
        self.row += 1
        for col in range(1, self.ncols + 1):
            self.ws.cell(self.row, col).fill = SECTION
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncols)
        self.cell(1, text, bold=True, size=12, color="92400E", fill=SECTION)
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1

    def kv(self, label, value, fmt=None, note=None, bold=False, span=2):
        """Etiqueta en A, valor en la columna `span`+1, explicación al lado."""
        self.cell(1, label, bold=bold, align=WRAP)
        if span > 1:
            self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=span)
        long_text = isinstance(value, str) and not note and span + 1 < self.ncols
        if long_text:  # un texto sin explicación ocupa el resto de la fila
            self.ws.merge_cells(start_row=self.row, start_column=span + 1, end_row=self.row, end_column=self.ncols)
        self.cell(span + 1, value, fmt=fmt, bold=bold, align=WRAP if isinstance(value, str) else None)
        if long_text and len(value) > 90:
            self.ws.row_dimensions[self.row].height = 15 * (1 + len(value) // 90)
        if note:
            self.ws.merge_cells(start_row=self.row, start_column=span + 2, end_row=self.row, end_column=self.ncols)
            self.cell(span + 2, note, color=MUTED, align=WRAP)
            if len(note) > 70:
                self.ws.row_dimensions[self.row].height = 15 * (1 + len(note) // 70)
        for col in range(1, self.ncols + 1):
            self.ws.cell(self.row, col).border = Border(bottom=LINE)
        self.row += 1

    def groups(self, spec):
        """Fila de grupos sobre el encabezado: [(desde, hasta, texto, clave de color), ...]."""
        for a, b, label, key in spec:
            fill, color = GROUPS[key]
            if b > a:
                self.ws.merge_cells(start_row=self.row, start_column=a, end_row=self.row, end_column=b)
            for col in range(a, b + 1):
                self.ws.cell(self.row, col).fill = fill
            self.cell(a, label, bold=True, color=color, fill=fill, align=CENTER)
        self.row += 1

    def header(self, labels, group_of=None):
        for i, label in enumerate(labels, 1):
            key = group_of(i) if group_of else None
            if key:
                fill, color = GROUPS[key]
                self.cell(i, label, bold=True, color=color, fill=fill, align=CENTER)
            else:
                self.cell(i, label, bold=True, color="FFFFFF", fill=HEAD, align=CENTER)
        self.ws.row_dimensions[self.row].height = 32
        self.row += 1
        return self.row

    def line(self, values, fmts, bold=False, fill=None, top=False):
        for i, v in enumerate(values, 1):
            fmt = fmts[i - 1] if i - 1 < len(fmts) else None
            if isinstance(v, F):
                v, fmt = v.value, v.fmt
            if v == "":
                v = None  # una celda vacía deja que el texto de al lado se vea completo
            c = self.cell(i, v, fmt=fmt, bold=bold, fill=fill)
            c.border = Border(bottom=LINE, top=STRONG if top else None)
            if isinstance(v, str) and len(v) > 28:
                c.alignment = WRAP
        self.row += 1

    def table(self, labels, rows, fmts, total=None, group_of=None, freeze=False, filters=True):
        first = self.header(labels, group_of)
        if freeze:
            self.ws.freeze_panes = self.ws.cell(first, 1)
        for r in rows:
            self.line(r, fmts)
        if filters and rows:
            self.ws.auto_filter.ref = f"A{first - 1}:{get_column_letter(len(labels))}{self.row - 1}"
        if total:
            self.line(total, fmts, bold=True, fill=SOFT, top=True)
        if not rows:
            self.text("Sin datos en este periodo.", color=MUTED, italic=True)


def check_row(sh, label, ok, detail):
    sh.cell(1, label, align=WRAP)
    sh.ws.merge_cells(start_row=sh.row, start_column=1, end_row=sh.row, end_column=2)
    sh.cell(3, "✔ Cuadra" if ok else "✖ Revisar", bold=True, color="166534" if ok else "991B1B",
            fill=OK_FILL if ok else BAD_FILL, align=CENTER)
    sh.ws.merge_cells(start_row=sh.row, start_column=4, end_row=sh.row, end_column=sh.ncols)
    sh.cell(4, detail, color=MUTED, align=WRAP)
    for col in range(1, sh.ncols + 1):
        sh.ws.cell(sh.row, col).border = Border(bottom=LINE)
    if len(detail) > 70:
        sh.ws.row_dimensions[sh.row].height = 15 * (1 + len(detail) // 70)
    sh.row += 1


def fmt_money(v):
    return "$ " + f"{v:,.0f}".replace(",", ".")


# ------------------------------------------------------------ hojas

def resumen(wb, c):
    s, f, pnl, year = c["summary"], c["summary"]["f210"]["casillas"], c["pnl"], c["year"]
    sh = Sheet(wb, "Resumen", [36, 18, 24, 22, 22, 22, 24], first=True)
    sh.title(f"Contabilidad de criptoactivos · Año gravable {year}",
             "Preparado con ArbiCrypto Wallet a partir de tus operaciones en Binance. Cifras en pesos colombianos (COP).")
    if c["in_progress"]:
        sh.text(f"El año {year} aún no termina: estas cifras van hasta el {c['generated']:%d/%m/%Y} y cambian con cada operación. "
                "Descarga este archivo otra vez después del 31 de diciembre para declarar.", color="92400E", bold=True)

    sh.section("Datos del contribuyente")
    sh.kv("Contribuyente", c["name"] or c["username"])
    sh.kv("Cédula / NIT", c["nit"] or "Complétalo en Ajustes")
    sh.kv("Periodo", f"01/01/{year} – 31/12/{year}")
    sh.kv("Método de costo", c["method_text"],
          note="Define qué compra se considera vendida primero. Se usa el mismo todo el año (art. 69 ET).")
    sh.kv("Tipo de actividad", "Habitual (negocio)" if c["habitual"] else "Ocasional",
          note="Compraventa habitual de cripto: las utilidades son renta ordinaria (cédula general)." if c["habitual"]
          else "Las cripto que tuviste 2 años o más pagan ganancia ocasional (15 %); el resto, renta ordinaria.")
    sh.kv("4x1000 (GMF)", c["gmf_text"])
    sh.kv("Generado", c["generated"].strftime("%d/%m/%Y %H:%M"))

    sh.section("1. Tu resultado del año")
    sh.text("Así se llega a tu ganancia o pérdida. Cada línea sale de las hojas que siguen.", color=MUTED)
    sh.kv("Ventas y cambios de cripto (valor de venta)", pnl["sales"], COP,
          note="Lo que recibiste en cada venta P2P y el valor comercial de cada cambio entre cripto (Concepto DIAN 1621 de 2023).")
    sh.kv("(−) Costo de lo vendido", -pnl["cost"], COP,
          note=f"Lo que te costó la cripto que vendiste, por {c['method_short']}. Ya incluye las comisiones de compra de Binance "
               f"({fmt_money(pnl.get('buy_fees', 0))}). Hoja Kardex.")
    if pnl["fees"]:
        sh.kv("(−) Comisiones de venta", -pnl["fees"], COP, note="Comisiones que Binance cobró al vender.")
    sh.kv("= Utilidad en operaciones", pnl["gross_profit"], COP, bold=True)
    if pnl["income"]:
        sh.kv("(+) Rendimientos (Earn) y airdrops", pnl["income"], COP, note="Valor de la cripto el día que la recibiste.")
    sh.kv("(−) 4x1000", -pnl["gmf"], COP, note="0,4 % de cada compra pagada desde tu banco ($4.000 por cada $1.000.000). Hoja 4x1000.")
    if pnl["expenses"]:
        sh.kv("(−) Otros gastos del negocio", -pnl["expenses"], COP)
    sh.kv("= Resultado neto del año", pnl["net"], COP, bold=True,
          note="Ganancia (positivo) o pérdida (negativo) de tu actividad con cripto en el año.")

    sh.section("2. Comisiones de Binance y 4x1000")
    fq = ", ".join(f"{q:.8g} {a}".replace(".", ",") for a, q in c["fee_qty"].items() if q > 0) or "—"
    sh.kv("Comisiones de Binance en compras", pnl.get("buy_fees", 0), COP,
          note=f"Binance se queda una parte de la cripto en cada orden P2P ({fq}). Te llega un poco menos de lo que pagaste; "
               "por eso esa comisión ya está dentro del costo.")
    sh.kv("Comisiones de Binance en ventas", pnl["fees"], COP)
    sh.kv("4x1000 pagado en el año", s["gmf"]["paid"], COP,
          note="Con certificado del banco: se usa el certificado." if s["gmf"]["cert"] else
          "Calculado sobre cada compra. Pide el certificado del GMF a tu banco para confirmarlo.")
    sh.kv("50 % del 4x1000 (deducible)", f.get("84", 0), COP, note="Va en la casilla 84 del formulario 210.")

    sh.section("3. Formulario 210 · qué va en cada casilla (solo la parte de cripto)")
    sh.header(["Casilla", "Valor", "Concepto en el formulario", "Qué significa", "", "", "Dónde verlo"])
    sh.ws.merge_cells(start_row=sh.row - 1, start_column=4, end_row=sh.row - 1, end_column=6)
    boxes = [
        ("29", "Patrimonio bruto", "Lo que te costaron las cripto que tenías el 31 de diciembre (costo fiscal, no precio de mercado). "
         "Súmale tus otros bienes: cuentas, casa, carro.", "Hoja Patrimonio 31 dic"),
        ("74", "Ingresos brutos · rentas no laborales", "Todo lo que vendiste o cambiaste a valor de venta, más rendimientos.", "Hojas Ganancias y Kardex"),
        ("77", "Costos y deducciones procedentes", "Lo que te costó lo que vendiste (con comisiones de Binance) y gastos del negocio.", "Hoja Kardex"),
        ("78", "Renta líquida", "Casilla 74 menos 77, si da positivo.", ""),
        ("84", "GMF · 50 % del 4x1000", "La mitad del 4x1000 que pagaste en el año.", "Hoja 4x1000"),
        ("86", "Rentas exentas y deducciones imputables", "El 4x1000 deducible, con tope del 40 % de los ingresos y 1.340 UVT. "
         "Solo aplica si hay renta líquida.", ""),
        ("87", "Renta líquida ordinaria", "Casilla 78 menos 86: lo que paga impuesto.", ""),
        ("88", "Pérdida líquida", "Si lo que te costó (77) supera lo que vendiste (74). Se compensa en los próximos años.", ""),
    ]
    if f.get("112", 0) > 0:
        boxes += [("112", "Ganancias ocasionales · ingresos", "Ventas de cripto que tenías hace 2 años o más.", "Hoja Ganancias"),
                  ("113", "Ganancias ocasionales · costos", "Lo que te costaron esas cripto.", ""),
                  ("115", "Ganancias ocasionales gravables", "112 − 113 − 114.", ""),
                  ("127", "Impuesto de ganancias ocasionales", "15 % de la casilla 115.", "")]
    for n, name, what, where in boxes:
        sh.cell(1, f"Casilla {n}", bold=True)
        sh.cell(2, f.get(n, 0.0), fmt=COP, bold=True)
        sh.cell(3, name, align=WRAP)
        sh.ws.merge_cells(start_row=sh.row, start_column=4, end_row=sh.row, end_column=6)
        sh.cell(4, what, color=MUTED, align=WRAP)
        sh.cell(7, where, color=MUTED, align=WRAP)
        for col in range(1, sh.ncols + 1):
            sh.ws.cell(sh.row, col).border = Border(bottom=LINE)
        sh.ws.row_dimensions[sh.row].height = 15 * max(2, 1 + len(what) // 60)
        sh.row += 1
    sh.kv("Impuesto estimado por las cripto", s["tax_total"], COP, bold=True,
          note="Lo que suben las cripto sobre tu impuesto, según la tabla del art. 241 ET. Tus otros ingresos (sueldo, etc.) cambian este valor.")
    tax_result = f.get("87", 0) - f.get("88", 0)
    if abs(tax_result - pnl["net"]) >= 1 and pnl["gmf"]:
        sh.skip()
        sh.text(f"¿Por qué la DIAN ve {fmt_money(tax_result)} y tu contabilidad {fmt_money(pnl['net'])}? Porque en la contabilidad se resta "
                f"todo el 4x1000 ({fmt_money(pnl['gmf'])}), pero para el impuesto solo se puede restar la mitad "
                f"({fmt_money(f.get('86', 0))}, art. 115 ET). La otra mitad es un gasto real que no es deducible.",
                color="92400E", bold=True)

    sh.section("4. Topes · ¿tienes que declarar?")
    names = {"patrimonio": ("Patrimonio al 31 de diciembre", "4.500 UVT"), "ingresos": ("Ingresos del año", "1.400 UVT"),
             "consignaciones": ("Consignaciones y depósitos", "1.400 UVT"), "compras": ("Compras y consumos", "1.400 UVT")}
    sh.header(["Tope", "Tu valor", "Límite", "¿Lo superas?", "Límite en UVT"])
    for k, (label, uvt) in names.items():
        t = s["topes"][k]
        sh.line([label, t["value"], t["limit"], "Sí: debes declarar" if t["over"] else "No", uvt],
                [None, COP, COP, None, None], bold=t["over"])
    sh.text(f"UVT {year}: {fmt_money(s['uvt'])}. Si superas cualquiera de los topes, declaras renta el año siguiente. "
            "Los topes incluyen todos tus ingresos y bienes, no solo las cripto.", color=MUTED)

    sh.section("5. Comprobaciones")
    for label, ok, detail in c["checks"]:
        check_row(sh, label, ok, detail)

    sh.section("6. Cómo leer este archivo")
    guide = [
        ("Kardex", "Libro de criptoactivos: por cada cripto, cada entrada y salida con su costo y el saldo corrido."),
        ("Operaciones", "Todas tus operaciones numeradas, con la orden de Binance, la comisión, el 4x1000 y el medio de pago."),
        ("Ganancias", "Cada venta o cambio: valor de venta, costo fiscal y utilidad. Es la base de las casillas 74 y 77."),
        ("Rondas", "Tus rondas de arbitraje: cuánto salió y entró de tus cuentas en cada vuelta."),
        ("4x1000", "El 4x1000 por mes y por cuenta, y la parte deducible."),
        ("Patrimonio 31 dic", "Las cripto que tenías al cierre del año y su costo fiscal (casilla 29)."),
        ("Libro diario", "Contabilidad por partida doble con cuentas del PUC: cada operación con su débito y crédito."),
        ("Mayor y balances", "Saldo de cada cuenta contable en el año."),
        ("Estados financieros", "Balance general y estado de resultados."),
    ]
    if c.get("banks"):
        guide.append(("Bancos", "Tus cuentas bancarias según los extractos que subiste."))
    for name, what in guide:
        sh.kv(name, what, span=1)
    sh.skip()
    sh.text("Cálculo informativo según el Estatuto Tributario y el Concepto Unificado DIAN 1621 de 2023: las cripto son activos "
            "intangibles; cada venta o cambio entre cripto es una enajenación a valor comercial y se declaran por su costo fiscal. "
            "Revísalo con tu contador antes de presentar la declaración.", color=MUTED, italic=True)


def kardex(wb, c):
    sh = Sheet(wb, "Kardex", [6, 18, 30, 29, 15, 16, 15, 15, 16, 15, 15, 15, 15, 14, 15, 15, 12])
    sh.title(f"Kardex de criptoactivos · {c['year']}",
             f"Método de costo: {c['method_text']}. Cada salida toma el costo de los lotes comprados. "
             "Utilidad = valor de venta − comisión − costo.")
    labels = ["N°", "Fecha", "Detalle", "Referencia", "Cantidad", "Costo unitario", "Costo total", "Cantidad", "Costo unitario",
              "Costo total", "Cantidad", "Costo total", "Valor de venta", "Utilidad", "Comisión Binance", "Comisión en cripto", "4x1000"]
    fmts = [None, DATE, None, None, QTY, COP, COP, QTY, COP, COP, QTY, COP, COP, COP, COP, QTY, COP]
    zone = lambda i: "in" if 5 <= i <= 7 else "out" if 8 <= i <= 10 else "bal" if 11 <= i <= 12 else "res" if 13 <= i <= 14 else "cost" if i >= 15 else None
    if not c["kardex"]:
        sh.text("Sin criptoactivos en este año.", color=MUTED, italic=True)
    for b in c["kardex"]:
        a, t = b["asset"], b["totals"]
        sh.skip()
        sh.cell(1, a, bold=True, size=13)
        sh.cell(3, f"{len(b['rows'])} movimientos · utilidad {fmt_money(t['gain'])}", color=MUTED)
        sh.row += 1
        sh.groups([(5, 7, "Entradas", "in"), (8, 10, "Salidas", "out"), (11, 12, "Saldo", "bal"),
                   (13, 14, "Resultado", "res"), (15, 17, "Costos de la operación", "cost")])
        sh.header(labels, zone)
        sh.line(["", f"01/01/{c['year']}", "Saldo inicial", "", "", "", "", "", "", "", b["opening"]["qty"], b["opening"]["cost"]],
                fmts, bold=True, fill=SOFT)
        for i, r in enumerate(b["rows"], 1):
            out = "out_qty" in r
            sh.line([i, xl_time(r["t"]), r["text"], r["ref"],
                     r.get("in_qty"), r.get("in_unit"), r.get("in_cost"),
                     r.get("out_qty"), r.get("out_unit"), r.get("out_cost"),
                     r["bal_qty"], r["bal_cost"],
                     r.get("sale") if out else None, r.get("gain") if out else None,
                     r.get("fee") or None, r.get("fee_qty") or None, r.get("gmf") or None], fmts)
        sh.line(["", "", "Totales y saldo final", "", t["in_qty"], None, t["in_cost"], t["out_qty"], None, t["out_cost"],
                 b["closing"]["qty"], b["closing"]["cost"], t["sales"], t["gain"], t.get("fees") or None, t.get("fee_qty") or None,
                 t.get("gmf") or None], fmts, bold=True, fill=SOFT, top=True)
        ok = abs(b["opening"]["cost"] + t["in_cost"] - t["out_cost"] - b["closing"]["cost"]) < 1
        sh.ws.merge_cells(start_row=sh.row, start_column=3, end_row=sh.row, end_column=9)
        sh.cell(3, f"Saldo inicial {fmt_money(b['opening']['cost'])} + entradas {fmt_money(t['in_cost'])} − salidas "
                   f"{fmt_money(t['out_cost'])} = saldo final {fmt_money(b['closing']['cost'])}", color=MUTED)
        sh.cell(10, "✔ Cuadra" if ok else "✖ Revisar", bold=True, color="166534" if ok else "991B1B", fill=OK_FILL if ok else BAD_FILL)
        sh.row += 1
    sh.skip()
    sh.text("La comisión de Binance en las compras ya está dentro del costo (te llega menos cripto de la que pagaste). "
            "El 4x1000 no es costo de la cripto: es un gasto aparte y la mitad es deducible (casilla 84).", color=MUTED)


def operaciones(wb, c):
    sh = Sheet(wb, "Operaciones", [15, 18, 20, 17, 29, 15, 7, 15, 7, 16, 16, 14, 13, 12, 16, 16, 17, 16, 24])
    sh.title(f"Registro de operaciones · {c['year']}",
             "Todas las operaciones del año, numeradas. Salió/Entró de tus cuentas es la plata real en tus bancos.")
    rows, tot = [], [0.0] * 6
    for m in c["ops"]:
        k = m["kind"]
        v, fee, gmf = float(m["cop_value"] or 0), float(m["fee_cop"] or 0), float(m["gmf_cop"] or 0)
        value = v + fee if k == "compra" else v
        gross_qty = (m["in_qty"] or 0) + (m.get("fee_qty") or 0) if k == "compra" else (m["out_qty"] or 0)
        unit = (value / gross_qty) if k in ("compra", "venta", "permuta") and gross_qty else (v / m["in_qty"] if k in ("ingreso", "airdrop") and m["in_qty"] else None)
        left = v + fee + gmf if k == "compra" else v + gmf if k in ("gasto", "retiro") else 0.0
        came = v - fee if k == "venta" else v if k == "aporte" else 0.0
        rows.append([m["n"], xl_time(m["t"]), KIND.get(k, k), m["platform"] or "", m.get("ref") or ("Manual" if m["source"] == "manual" else ""),
                     F(m["out_qty"], COP) if m["out_asset"] == "COP" else m["out_qty"], m["out_asset"] or "",
                     F(m["in_qty"], COP) if m["in_asset"] == "COP" else m["in_qty"], m["in_asset"] or "", unit, value, fee or None,
                     m.get("fee_qty") or None, gmf or None, left or None, came or None, m.get("pay_label") or "",
                     m["counterpart"] or "", m["note"] or ""])
        for i, x in enumerate((value, fee, gmf, left, came)):
            tot[i] += x
    fmts = [None, DATE, None, None, None, QTY, None, QTY, None, COP, COP, COP, QTY, COP, COP, COP, None, None, None]
    sh.table(["N°", "Fecha", "Tipo", "Plataforma", "Referencia (orden de Binance)", "Entregaste", "", "Recibiste", "",
              "Precio unitario", "Valor de la operación", "Comisión Binance", "Comisión en cripto", "4x1000",
              "Salió de tus cuentas", "Entró a tus cuentas", "Medio de pago", "Contraparte", "Nota"], rows, fmts,
             total=[f"Totales ({len(rows)})", "", "", "", "", None, "", None, "", None, tot[0], tot[1], None, tot[2], tot[3], tot[4], "", "", ""],
             freeze=True)
    sh.skip()
    sh.text("Valor de la operación: en compras, lo que pagaste (incluye la comisión de Binance); en ventas, lo que vendiste; "
            "en cambios entre cripto, el valor comercial de ese momento.", color=MUTED)


def ganancias(wb, c):
    sh = Sheet(wb, "Ganancias", [18, 20, 9, 16, 10, 16, 14, 16, 16, 16, 16, 12, 13])
    sh.title(f"Ganancias por operación · {c['year']}",
             "Cada venta o cambio de cripto es una enajenación (art. 90 ET y Concepto DIAN 1621 de 2023). Base de las casillas 74 y 77.")
    rows = [[xl_time(x["t"]), KIND.get(x["kind"], x["kind"]), x["asset"], x["qty"], x.get("received") or "", x["gross"], x["fee"] or None,
             x["cost"], x["gain"], x["ro_gain"], x["go_gain"] or None, round(x["days_max"], 1), "Sí" if x["missing"] > 1e-9 else ""]
            for x in c["disposals"]]
    rows += [[xl_time(x["t"]), KIND.get(x["kind"], x["kind"]), x["asset"], x["qty"], "", x["value"], None, 0.0, x["value"],
              x["value"] if x["kind"] == "ingreso" else 0.0, x["value"] if x["kind"] == "airdrop" else None, None, ""]
             for x in c["incomes"]]
    rows.sort(key=lambda r: r[0])
    tot = lambda i: sum((r[i] or 0) for r in rows)
    sh.table(["Fecha", "Tipo", "Cripto", "Cantidad", "Recibiste", "Valor de venta", "Comisión", "Costo fiscal", "Utilidad",
              "Renta ordinaria", "Ganancia ocasional", "Días de tenencia", "Sin compra registrada"], rows,
             [DATE, None, None, QTY, None, COP, COP, COP, COP, COP, COP, "0.0", None],
             total=["Totales", "", "", None, "", tot(5), tot(6), tot(7), tot(8), tot(9), tot(10), None, ""], freeze=True)


def rondas(wb, c):
    sh = Sheet(wb, "Rondas", [6, 18, 18, 12, 24, 16, 16, 14, 16, 16, 12])
    sh.title(f"Rondas de arbitraje · {c['year']}",
             "Una ronda va desde que sale plata a comprar cripto hasta que todo vuelve a pesos. Resultado = recibiste − pagaste − 4x1000.")
    rows = [[i, xl_time(r["start"]), xl_time(r["end"]), r["ops"], " → ".join(r["assets"]), r["paid"], r.get("binance_fees") or None,
             r["gmf"] or None, r["received"], r["result"], "Abierta" if r["open"] else "Cerrada"] for i, r in enumerate(c["rounds"], 1)]
    tot = lambda i: sum((r[i] or 0) for r in rows)
    sh.table(["N°", "Inicio", "Fin", "Operaciones", "Cripto", "Pagaste", "Comisión Binance (incluida)", "4x1000", "Recibiste",
              "Resultado", "Estado"], rows, [None, DATE, DATE, None, None, COP, COP, COP, COP, COP, None],
             total=["", "Totales", "", tot(3), "", tot(5), tot(6), tot(7), tot(8), tot(9), ""], freeze=True)
    sh.skip()
    sh.text("La comisión de Binance ya va dentro de lo pagado (y de lo recibido, en ventas). Se muestra para que sepas cuánto se quedó Binance.",
            color=MUTED)


def gmf(wb, c):
    s = c["summary"]
    sh = Sheet(wb, "4x1000", [30, 20, 16, 16, 40])
    sh.title(f"4x1000 (GMF) · {c['year']}", "El banco cobra 0,4 % ($4.000 por cada $1.000.000) de lo que sale de tu cuenta.")
    sh.kv("Cómo lo cobra tu banco", c["gmf_text"], span=1)
    months = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    sh.section("Por mes")
    rows = [[months[m["month"] - 1], m["bought"], m["gmf"], m["ops"]] for m in c["months"] if m["bought"] or m["gmf"]]
    sh.table(["Mes", "Pagaste por cripto", "4x1000", "Operaciones del mes"], rows, [None, COP, COP, None],
             total=["Total", sum(r[1] for r in rows), sum(r[2] for r in rows), sum(r[3] for r in rows)], filters=False)
    sh.section("Por cuenta o medio de pago")
    rows = [[r["method"], r["out"], r["gmf"], r["ops"]] for r in c["by_method"] if r["out"] or r["gmf"]]
    sh.table(["Medio de pago", "Salió de la cuenta", "4x1000", "Pagos"], rows, [None, COP, COP, None],
             total=["Total", sum(r[1] for r in rows), sum(r[2] for r in rows), sum(r[3] for r in rows)], filters=False)
    sh.section("En la declaración")
    sh.kv("4x1000 calculado", s["gmf"]["estimated"], COP, span=1)
    if s["gmf"]["cert"]:
        sh.kv("4x1000 según certificado o extractos", s["gmf"]["cert"], COP, span=1, note="Manda sobre el calculado.")
    sh.kv("4x1000 pagado", s["gmf"]["paid"], COP, span=1, bold=True)
    sh.kv("Casilla 84 · 50 % deducible", s["f210"]["casillas"].get("84", 0), COP, span=1)
    sh.kv("Casilla 86 · lo que se puede restar", s["f210"]["casillas"].get("86", 0), COP, span=1,
          note="Tope: 40 % de los ingresos y 1.340 UVT, y solo si hay renta líquida (no crea ni agranda pérdidas).")


def patrimonio(wb, c):
    sh = Sheet(wb, "Patrimonio 31 dic", [12, 20, 18, 18, 17])
    sh.title(f"Criptoactivos al 31 de diciembre de {c['year']}",
             "Se declaran por su costo fiscal (lo que te costaron), no por el precio de mercado. Va en la casilla 29.")
    if c["in_progress"]:
        sh.text(f"El año aún no termina: esto es lo que tienes hoy ({c['generated']:%d/%m/%Y}). El 31 de diciembre se actualiza solo.",
                color="92400E")
    rows = [[h["asset"], h["qty"], h["cost"], h["cost"] / h["qty"] if h["qty"] else None, xl_time(h["since"])] for h in c["holdings"]]
    if rows:
        sh.table(["Cripto", "Cantidad", "Costo fiscal", "Costo unitario", "Lote más antiguo"], rows, [None, QTY, COP, COP, DAY],
                 total=["Total", None, sum(r[2] for r in rows), None, None], filters=False)
    else:
        sh.text("No tienes cripto: todo está en pesos. La casilla 29 por cripto es $0.", bold=True)


def diario(wb, c):
    sh = Sheet(wb, "Libro diario", [14, 17, 44, 10, 34, 16, 16])
    sh.title(f"Libro diario · {c['year']}", "Partida doble con cuentas del PUC. En cada comprobante, débitos = créditos.")
    rows = [[e["number"], xl_time(e["t"]), e["text"], ln["account"], ln["name"], ln["debit"] or None, ln["credit"] or None]
            for e in c["journal"] for ln in e["lines"]]
    sh.table(["Comprobante", "Fecha", "Detalle", "Cuenta", "Nombre de la cuenta", "Débito", "Crédito"], rows,
             [None, DATE, None, None, None, COP, COP],
             total=["Totales", "", "", "", "", sum(r[5] or 0 for r in rows), sum(r[6] or 0 for r in rows)], freeze=True)


def mayor(wb, c):
    tb = c["trial"]
    sh = Sheet(wb, "Mayor y balances", [10, 40, 18, 18, 18, 18])
    sh.title(f"Libro mayor y balances · {c['year']}", "Saldo de cada cuenta: saldo anterior + débitos − créditos = saldo final.")
    rows = [[r["account"], r["name"], r["previous"], r["debit"], r["credit"], r["final"]] for r in tb["rows"]]
    sh.table(["Cuenta", "Nombre", "Saldo anterior", "Débitos", "Créditos", "Saldo final"], rows, [None, None, COP, COP, COP, COP],
             total=["", "Totales", None, tb["totals"]["debit"], tb["totals"]["credit"], None], freeze=True)


def estados(wb, c):
    bs, inc = c["balance"], c["income"]
    sh = Sheet(wb, "Estados financieros", [10, 46, 20, 40])
    sh.title(f"Estados financieros · {c['year']}", f"Balance general al {c['at']:%d/%m/%Y} y estado de resultados del año.")
    sh.section("Balance general")
    sh.header(["Cuenta", "Activo", "Valor"])
    for x in bs["assets"]:
        sh.line([x["account"], x["name"], x["value"]], [None, None, COP])
    sh.line(["", "Total activo", bs["total_assets"]], [None, None, COP], bold=True, fill=SOFT, top=True)
    sh.skip()
    sh.header(["Cuenta", "Pasivo", "Valor"])
    for x in bs["liabilities"]:
        sh.line([x["account"], x["name"], x["value"]], [None, None, COP])
    sh.line(["", "Total pasivo", bs["total_liabilities"]], [None, None, COP], bold=True, fill=SOFT, top=True)
    sh.skip()
    sh.header(["Cuenta", "Patrimonio", "Valor"])
    for x in bs["equity"]:
        sh.line([x["account"], x["name"], x["value"]], [None, None, COP])
    sh.line(["", "Total patrimonio", bs["total_equity"]], [None, None, COP], bold=True, fill=SOFT, top=True)
    sh.line(["", "Total pasivo y patrimonio", bs["total_liabilities"] + bs["total_equity"]], [None, None, COP], bold=True, fill=SOFT, top=True)
    sh.section("Estado de resultados")
    sh.header(["Cuenta", "Concepto", "Valor"])
    for x in inc["layout"]:
        if x["kind"] == "section":
            sh.line(["", x["label"], ""], [None, None, COP], bold=True)
        else:
            strong = x["kind"] != "line"
            sh.line([x["account"], x["label"], x["value"]], [None, None, COP], bold=strong, fill=SOFT if strong else None, top=x["kind"] == "net")
    cash = sum(x["value"] for x in bs["assets"] if x["account"] == "1110")
    if cash < -1:
        sh.skip()
        sh.text(f"Bancos queda en {fmt_money(cash)} porque los libros no saben con cuánta plata empezaste. Regístralo en la app "
                "(Inicio → Inicia tu caja). No cambia la declaración de las cripto.", color="92400E")


def patrimonio_cambios(wb, c):
    eq = c["equity"]
    cols = eq["columns"]
    sh = Sheet(wb, "Cambios en el patrimonio", [40, *[20] * len(cols), 20])
    sh.title(f"Estado de cambios en el patrimonio · {c['year']}", "Cómo pasó el patrimonio del saldo inicial al final.")
    sh.header(["Concepto", *[f"{x['account']} {x['name']}" for x in cols], "Total"])
    for r in eq["rows"]:
        strong = r["key"] in ("opening", "closing")
        sh.line([r["label"], *[r["values"][x["account"]] for x in cols], r["total"]], [None, *[COP] * (len(cols) + 1)],
                bold=strong, fill=SOFT if strong else None, top=r["key"] == "closing")


def flujo(wb, c):
    cf = c["cashflow"]
    sh = Sheet(wb, "Flujo de efectivo", [52, 20])
    sh.title(f"Estado de flujos de efectivo · {c['year']}", "Método directo: de dónde vino y a dónde se fue el dinero de caja y bancos.")
    sh.header(["Concepto", "Valor"])
    sh.line(["Efectivo al inicio del periodo", cf["opening"]], [None, COP], bold=True)
    for s in cf["sections"]:
        sh.line([s["label"], ""], [None, COP], bold=True)
        for ln in s["lines"]:
            sh.line(["   " + ln["label"], ln["value"]], [None, COP])
        sh.line([f"Efectivo neto de {s['label'].lower()}", s["total"]], [None, COP], bold=True, fill=SOFT)
    sh.line(["Aumento (disminución) neto del efectivo", cf["net"]], [None, COP], bold=True, top=True)
    sh.line(["Efectivo al final del periodo", cf["closing"]], [None, COP], bold=True, fill=SOFT)


def inventarios(wb, c):
    inv = c["inventory"]
    b = inv["balance"]
    sh = Sheet(wb, "Inventarios y balances", [12, 40, 18, 18, 20, 20])
    sh.title(f"Libro de inventarios y balances · corte {c['at']:%d/%m/%Y}",
             "Cada bien, deuda y cuenta de patrimonio al corte. Las cripto, una por una, al costo; el valor de mercado es informativo.")
    fmts = [None, None, QTY, COP, COP, COP]
    sh.section("Activo")
    sh.header(["Cuenta", "Nombre", "Cantidad", "Costo unitario", "Valor en libros", "Valor de mercado"])
    for x in b["assets"]:
        sh.line([x["account"], x["name"], "", "", x["value"], ""], fmts, bold=x["account"] == inv["account"] and bool(inv["cryptos"]))
        if x["account"] == inv["account"]:
            for k in inv["cryptos"]:
                sh.line(["", "   " + k["asset"], k["qty"], k["unit"], k["cost"], k["market"] if k["market"] is not None else ""], fmts)
    sh.line(["", "Total activo", "", "", b["total_assets"], ""], fmts, bold=True, fill=SOFT, top=True)
    sh.section("Pasivo")
    for x in b["liabilities"]:
        sh.line([x["account"], x["name"], "", "", x["value"], ""], fmts)
    sh.line(["", "Total pasivo", "", "", b["total_liabilities"], ""], fmts, bold=True, fill=SOFT, top=True)
    sh.section("Patrimonio")
    for x in b["equity"]:
        sh.line([x["account"], x["name"], "", "", x["value"], ""], fmts)
    sh.line(["", "Total patrimonio", "", "", b["total_equity"], ""], fmts, bold=True, fill=SOFT, top=True)
    sh.line(["", "Total pasivo y patrimonio", "", "", b["total_liabilities"] + b["total_equity"], ""], fmts, bold=True, fill=SOFT)
    sh.skip(2)
    sh.text("_______________________________          _______________________________")
    sh.text("Contribuyente                                              Contador público · T.P.")


def notas(wb, c):
    sh = Sheet(wb, "Notas", [26, 18, 18, 18, 18, 18])
    sh.title(f"Notas a los estados financieros · {c['year']}", f"{c['name'] or c['username']} · cédula {c['nit'] or '—'}")
    for n in c["notes"]:
        sh.section(f"Nota {n['n']}. {n['title']}")
        for para in n["text"]:
            sh.text(para)
        t = n.get("table")
        if t and t["rows"]:
            sh.skip()
            fmts = [None if i == 0 or any(isinstance(r[i], str) for r in t["rows"]) else (QTY if t["head"][i] == "Cantidad" else
                    None if t["head"][i] == "Operaciones" else COP) for i in range(len(t["head"]))]
            sh.table(t["head"], [[("" if v is None else v) for v in r] for r in t["rows"]], fmts, filters=False)


def caja(wb, c):
    b = c["cash_book"]
    sh = Sheet(wb, "Caja diaria", [14, 17, 44, 16, 16, 14, 18])
    sh.title(f"Libro de caja y bancos · cierre diario · {c['year']}",
             "Cuenta 1110 Bancos. Cada día se cierra solo: saldo inicial + entradas − salidas − 4x1000 = saldo final.")
    fmts = [None, DATE, None, COP, COP, COP, COP]
    first = sh.header(["Comprobante", "Fecha", "Detalle", "Entradas", "Salidas", "4x1000", "Saldo"])
    sh.ws.freeze_panes = sh.ws.cell(first, 1)
    sh.line(["", "", "Saldo inicial", "", "", "", b["opening"]], fmts, bold=True)
    for d in b["days"]:
        for e in d["entries"]:
            sh.line([e["number"], xl_time(e["t"]), e["text"], e["in"] or "", e["out"] or "", e["gmf"] or "", e["balance"]], fmts)
        day = f"{d['day'][8:10]}/{d['day'][5:7]}/{d['day'][:4]}"
        sh.line(["", "", f"Cierre del día {day}" + (" (abierto)" if d["status"] == "abierto" else ""), d["in"], d["out"], d["gmf"] or "",
                 d["closing"]], fmts, bold=True, fill=SOFT)
    t = b["totals"]
    sh.line(["Totales", "", "Saldo final", t["in"], t["out"], t["gmf"], b["closing"]], fmts, bold=True, fill=SOFT, top=True)


def bancos(wb, c):
    b = c["banks"]
    sh = Sheet(wb, "Bancos", [28, 14, 18, 18, 16, 18, 14])
    sh.title(f"Cuentas bancarias · {c['year']}", "Según los extractos que subiste.")
    rows = [[a["name"], a.get("kind") or "", a["in"], a["out"], a["gmf"] or None, a["balance_end"], a["n"]] for a in b["accounts"]]
    sh.table(["Cuenta", "Tipo", "Entradas", "Salidas", "4x1000 real", "Saldo al cierre", "Movimientos"], rows,
             [None, None, COP, COP, COP, COP, None], filters=False)
    d = b["dian"]
    sh.section("Para tu declaración")
    sh.kv("Saldo de tus cuentas al 31 de diciembre", d["bank_balances_end"], COP, note="Va en el patrimonio (casilla 29) con tus otros bienes.")
    sh.kv("Consignaciones que no son de cripto", d["deposits_other"], COP, note="Cuentan para el tope de consignaciones.")
    if d.get("salary"):
        sh.kv("Salario", d["salary"], COP)
    if d.get("other_income"):
        sh.kv("Otros ingresos y rendimientos", d["other_income"], COP)
    sh.kv("Compras y consumos", d["consumption"], COP, note="Cuentan para el tope de compras y consumos.")
    sh.kv("4x1000 según extractos", d["gmf_real"], COP)


def cierres(wb, c):
    sh = Sheet(wb, "Cierres", [17, 12, 18, 18, 20, 11, 30])
    sh.title("Cierres de caja", "Cada cierre compara la caja según los libros con tu saldo real; la diferencia entra como aporte o retiro.")
    kinds = {"caja": "Caja", "inicio": "Inicio de caja", "mes": "Mes", "anio": "Año"}
    rows = [[xl_time(x["t"]), kinds.get(x["kind"], x["kind"]), x["cash_expected"], x["cash_real"], x["adjustment"],
             "Sí" if x["locked"] else "No", x["note"]] for x in c["closings"]]
    sh.table(["Fecha", "Tipo", "Caja según libros", "Caja real", "Diferencia", "Congelado", "Nota"], rows,
             [DATE, None, COP, COP, COP, None, None], filters=False)


def build(c):
    wb = Workbook()
    resumen(wb, c)
    kardex(wb, c)
    operaciones(wb, c)
    ganancias(wb, c)
    rondas(wb, c)
    gmf(wb, c)
    patrimonio(wb, c)
    diario(wb, c)
    mayor(wb, c)
    inventarios(wb, c)
    estados(wb, c)
    patrimonio_cambios(wb, c)
    flujo(wb, c)
    notas(wb, c)
    if c["cash_book"]["days"]:
        caja(wb, c)
    if c.get("banks"):
        bancos(wb, c)
    if c["closings"]:
        cierres(wb, c)
    wb.properties.title = f"Contabilidad de criptoactivos {c['year']}"
    wb.properties.creator = "ArbiCrypto Wallet"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
