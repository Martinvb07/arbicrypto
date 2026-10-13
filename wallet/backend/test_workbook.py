"""Pruebas del libro de Excel para la DIAN:  python -m unittest test_workbook -v"""
import io
import unittest
from datetime import datetime

from openpyxl import load_workbook

import ledger
import tax
import workbook
from test_tax import D, buy, mv, sell


def context(ms, year=2026):
    start, end = tax.year_bounds(year)
    res = tax.compute(ms, year, True, "fifo")
    s = tax.summary(res, year)
    bk = tax.books(res, ms, year)
    entries = ledger.journal(ms, [], True, "fifo")
    at = end - 1
    ops = [{**m, "n": f"OP-{year}-{i:04d}", "ref": "", "pay_label": "Nequi", "platform": "", "counterpart": "", "source": "binance",
            "note": "", "fee_qty": m.get("fee_qty", 0.0)} for i, m in enumerate(ms, 1)]
    return {
        "year": year, "generated": datetime(2027, 1, 5, tzinfo=tax.CO_TZ), "in_progress": False, "at": datetime.fromtimestamp(at, tax.CO_TZ),
        "name": "Prueba", "nit": "123", "username": "p", "habitual": True, "method_text": "PEPS", "method_short": "PEPS",
        "gmf_text": "Se paga en todas las compras.", "summary": s, "pnl": bk["pnl"], "months": bk["months"], "rounds": bk["rounds"],
        "kardex": tax.kardex(ms, year, "fifo"), "ops": ops, "fee_qty": {"USDT": 0.07}, "disposals": res["disposals"],
        "incomes": res["incomes"], "holdings": [{"asset": a, **h} for a, h in res["holdings"].items()],
        "journal": [{**e, "lines": [{**ln, "name": ledger.ACCOUNTS[ln["account"]][0]} for ln in e["lines"]]} for e in entries if start <= e["t"] < end],
        "trial": ledger.trial_balance(entries, start, end), "balance": ledger.balance_sheet(entries, at),
        "income": ledger.income_statement(entries, start, end), "account_name": lambda c: ledger.ACCOUNTS[c][0],
        "closings": [], "banks": None, "by_method": [{"method": "Nequi", "out": 400_000, "gmf": 1600, "ops": 1}],
        "checks": [("Kardex de cada cripto", True, "Cuadra.")],
        "cash_book": {**(cb := ledger.daily_cash(entries, start, end)),
                      "days": [{**d, "status": "cerrado", "real": None} for d in cb["days"]]},
        "inventory": {"at": at, "account": "1435", "balance": ledger.balance_sheet(entries, at),
                      "cryptos": [{"asset": a, "qty": h["qty"], "cost": h["cost"], "unit": h["cost"] / h["qty"], "since": h["since"],
                                   "market": None, "usd": None, "trm": None} for a, h in res["holdings"].items() if h["qty"] > 0]},
        "equity": ledger.equity_changes(entries, start, end), "cashflow": ledger.cash_flow(entries, start, end),
        "notes": [{"n": 1, "title": "Entidad que reporta", "text": ["Prueba."], "table": {"head": ["Cripto", "Cantidad", "Valor"], "rows": [["BTC", 0.1, 100.0]]}}],
    }


class Libro(unittest.TestCase):
    def test_libro_completo(self):
        compra = {**buy(D(2026, 3, 1), "USDT", 99.93, 399_720), "fee_cop": 280.0, "fee_qty": 0.07, "fee_asset": "USDT", "gmf_cop": 1600.0}
        ms = [compra, mv(D(2026, 3, 2), "permuta", "USDT", 99.93, "BTC", 0.001, cop=405_000), sell(D(2026, 3, 3), "BTC", 0.001, 410_000)]
        wb = load_workbook(io.BytesIO(workbook.build(context(ms))))
        self.assertEqual(wb.sheetnames[:7], ["Resumen", "Kardex", "Operaciones", "Ganancias", "Rondas", "4x1000", "Patrimonio 31 dic"])
        res = wb["Resumen"]
        found = {r[0].value: r[1].value for r in res.iter_rows(min_col=1, max_col=2) if isinstance(r[0].value, str)}
        self.assertAlmostEqual(found["Casilla 74"], 405_000 + 410_000)
        self.assertAlmostEqual(found["Casilla 84"], 800)           # la mitad del 4x1000
        # la comisión de Binance aparece en el kardex y va dentro del costo
        usdt = next(r for r in wb["Kardex"].iter_rows(values_only=True) if r[2] == "Totales y saldo final" and r[4] == 99.93)
        self.assertAlmostEqual(usdt[6], 400_000)                   # costo total de entradas USDT = lo pagado
        self.assertAlmostEqual(usdt[14], 280)                      # comisión Binance
        self.assertAlmostEqual(usdt[16], 1600)                     # 4x1000
        # en operaciones, lo que salió del banco = pagado + 4x1000
        op = next(r for r in wb["Operaciones"].iter_rows(values_only=True) if r[0] == "OP-2026-0001")
        self.assertAlmostEqual(op[10], 400_000)
        self.assertAlmostEqual(op[14], 401_600)
        # caja diaria: un cierre por día con movimiento en bancos y el saldo final = Bancos
        caja = list(wb["Caja diaria"].iter_rows(values_only=True))
        self.assertEqual(sum(1 for r in caja if str(r[2]).startswith("Cierre del día")), 2)  # 1 mar (compra) y 3 mar (venta)
        final = next(r for r in caja if r[0] == "Totales")
        self.assertAlmostEqual(final[6], -401_600 + 410_000)
        # libros oficiales y estados financieros completos, en orden
        books = ["Libro diario", "Mayor y balances", "Inventarios y balances", "Estados financieros", "Cambios en el patrimonio",
                 "Flujo de efectivo", "Notas", "Caja diaria"]
        self.assertEqual([s for s in wb.sheetnames if s in books], books)
        flujo = {r[0]: r[1] for r in wb["Flujo de efectivo"].iter_rows(values_only=True) if r[0]}
        self.assertAlmostEqual(flujo["Efectivo al final del periodo"], final[6])
        er = {r[1]: r[2] for r in wb["Estados financieros"].iter_rows(values_only=True) if r[1]}
        self.assertIn("Utilidad bruta", er)


if __name__ == "__main__":
    unittest.main()
